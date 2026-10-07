"""Izvoz prezentacije "Python za biologe" po lekcijama u PDF i PPTX.

Izvor su HTML slajdovi u ../izvor (deck.json + slides/*.html), u formatu
Claude Slides artefakta. Skripta:

1. deli slajdove na lekcije prema sekcijama iz deck.json (jedna lekcija po
   poglavlju knjige; uvodni slajdovi idu u lekciju 1, završni u lekciju 9),
2. PDF pravi tako što slajdove iscrta u Chromiumu (Playwright), sa
   originalnim fontovima IBM Plex Sans i IBM Plex Mono (Google Fonts),
3. PPTX pravi kao izmenljiv fajl: raspored elemenata meri u Chromiumu sa
   fontovima Arial i Consolas (koji postoje uz Office), pa svaki tekst,
   okvir i strelicu prenosi kao zaseban PowerPoint objekat; beleške iz
   <aside> postaju beleške predavača.

Pokretanje (iz bilo kog foldera):
    python prezentacija/alati/izvoz.py            # sve lekcije
    python prezentacija/alati/izvoz.py 3 5        # samo lekcije 3 i 5

Potrebno: python-pptx, playwright (python -m playwright install chromium)
i internet veza za Google Fonts (samo za PDF).
"""

import json
import re
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
IZVOR = ROOT / "izvor"
LEKCIJE = ROOT / "lekcije"

# naziv fajla i naslov lekcije, redom kao sekcije p1..p9 u deck.json
LESSONS = [
    ("01-uvod-i-radno-okruzenje", "Uvod i radno okruženje"),
    ("02-stampanje-i-rad-sa-tekstom", "Štampanje i rad sa tekstom"),
    ("03-citanje-i-pisanje-fajlova", "Čitanje i pisanje fajlova"),
    ("04-liste-i-petlje", "Liste i petlje"),
    ("05-pisanje-sopstvenih-funkcija", "Pisanje sopstvenih funkcija"),
    ("06-uslovni-testovi", "Uslovni testovi"),
    ("07-regularni-izrazi", "Regularni izrazi"),
    ("08-recnici", "Rečnici"),
    ("09-fajlovi-programi-i-korisnicki-unos", "Fajlovi, programi i korisnički unos"),
]

COURSE = "Python za biologe"
AUTHOR = "Aleksandar Kartelj"
PPTX_SANS = "Arial"
PPTX_MONO = "Consolas"

# slajd je 1920x1080 px, a PPTX 13,333 x 7,5 in: 1 px = 6350 EMU = 0,5 pt
EMU_PER_PX = 6350
PT_PER_PX = 0.5

FONTS_LINK = (
    '<link rel="stylesheet" href="https://fonts.googleapis.com/css2?'
    "family=IBM+Plex+Sans:ital,wght@0,400;0,500;0,600;0,700;1,400"
    '&family=IBM+Plex+Mono:wght@400;600&display=block">'
)

BASE_CSS = """
@page { size: 1920px 1080px; margin: 0; }
html, body { margin: 0; padding: 0; background: #ffffff; }
* { margin: 0; padding: 0; box-sizing: border-box; }
section { width: 1920px; height: 1080px; position: relative; overflow: hidden;
          display: flex; flex-direction: column;
          break-after: page; page-break-after: always; }
div:not([style*="display"]) { display: flex; flex-direction: column; }
div { min-width: 0; }
aside { display: none; }
h1 { font-size: 96px; font-weight: 600; line-height: 1.1; }
h2 { font-size: 64px; font-weight: 600; line-height: 1.15; }
h3 { font-size: 44px; font-weight: 600; line-height: 1.2; }
p { font-size: 32px; line-height: 1.4; }
section, section * { font-variant-ligatures: none; font-feature-settings: "liga" 0, "calt" 0; }
table { border-collapse: collapse; }
th, td { padding: 0.35em 0.6em; border-bottom: 1px solid #DCDDD5;
         text-align: left; vertical-align: top; }
th { font-weight: 600; }
x-shape, x-icon { display: block; flex: none; }
x-shape[kind="arrow-right"] {
  clip-path: polygon(0 30%, 60% 30%, 60% 0, 100% 50%, 60% 100%, 60% 70%, 0 70%); }
x-icon svg { width: 100%; height: 100%; display: block; }
"""

# jednostavne ikonice (linijski crtež, viewBox 24x24) za <x-icon name="...">
ICONS = {
    "Lightbulb": '<path d="M15 14c.2-1 .7-1.7 1.5-2.5 1-.9 1.5-2.2 1.5-3.5A6 6 0 0 0 6 8c0 1 .2 2.2 1.5 3.5.7.7 1.3 1.5 1.5 2.5"/><path d="M9 18h6"/><path d="M10 22h4"/>',
    "Tool": '<path d="M14.7 6.3a1 1 0 0 0 0 1.4l1.6 1.6a1 1 0 0 0 1.4 0l3.77-3.77a6 6 0 0 1-7.94 7.94l-6.91 6.91a2.12 2.12 0 0 1-3-3l6.91-6.91a6 6 0 0 1 7.94-7.94l-3.76 3.76z"/>',
    "Lightning": '<path d="M13 2 3 14h9l-1 8 10-12h-9l1-8z"/>',
    "Clock": '<circle cx="12" cy="12" r="10"/><polyline points="12 6 12 12 16 14"/>',
    "Warning": '<path d="m21.73 18-8-14a2 2 0 0 0-3.48 0l-8 14A2 2 0 0 0 4 21h16a2 2 0 0 0 1.73-3"/><path d="M12 9v4"/><path d="M12 17h.01"/>',
}

ICON_SCRIPT = """
(icons) => {
  document.querySelectorAll('x-icon').forEach((el, i) => {
    const body = icons[el.getAttribute('name')] || '<circle cx="12" cy="12" r="9"/>';
    el.innerHTML = '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" ' +
                   'stroke-linecap="round" stroke-linejoin="round">' + body + '</svg>';
    el.setAttribute('data-icon-id', String(i));
  });
}
"""

# za PPTX: označi monospace elemente, pa ceo slajd prebaci na Arial/Consolas
MEASURE_FONTS_SCRIPT = """
([sans, mono]) => {
  document.querySelectorAll('section *').forEach(el => {
    if (getComputedStyle(el).fontFamily.includes('Plex Mono')) el.setAttribute('data-mono', '');
  });
  const st = document.createElement('style');
  st.textContent = `section [data-mono] { font-family: '${mono}', monospace !important; }
                    section :not([data-mono]) { font-family: '${sans}', sans-serif !important; }`;
  document.head.appendChild(st);
}
"""

# provera rasporeda: deca koja izlaze iz roditelja ili sa slajda
OVERFLOW_SCRIPT = """
() => {
  const out = [];
  document.querySelectorAll('section').forEach(sec => {
    const S = sec.getBoundingClientRect();
    sec.querySelectorAll('*').forEach(el => {
      if (el.closest('aside') || el.closest('x-icon')) return;
      const r = el.getBoundingClientRect();
      if (r.width === 0 && r.height === 0) return;
      const issues = [];
      if (r.bottom - S.top > 1080 - 24) issues.push('dno slajda');
      if (r.right - S.left > 1920 - 24) issues.push('desna ivica slajda');
      const p = el.parentElement;
      if (p && p !== sec && getComputedStyle(el).position !== 'absolute') {
        const R = p.getBoundingClientRect();
        if (r.right > R.right + 2) issues.push('šire od roditelja');
        if (r.bottom > R.bottom + 2) issues.push('niže od roditelja');
      }
      if (['P','H1','H2','H3','TD','TH','LI'].includes(el.tagName) && el.scrollWidth > el.clientWidth + 2)
        issues.push('tekst šireg sadržaja');
      if (issues.length) out.push({slide: sec.id, tag: el.tagName,
        text: (el.textContent || '').trim().slice(0, 50), issues});
    });
  });
  return out;
}
"""

EXTRACT_SCRIPT = r"""
() => {
  const px = v => parseFloat(v) || 0;
  const color = c => {
    const m = /rgba?\(([^)]+)\)/.exec(c || '');
    if (!m) return null;
    const p = m[1].split(',').map(s => parseFloat(s));
    if (p.length > 3 && p[3] === 0) return null;
    return p.slice(0, 3).map(x => Math.round(x).toString(16).padStart(2, '0')).join('').toUpperCase();
  };
  const TEXT = new Set(['P', 'H1', 'H2', 'H3', 'TH', 'TD']);
  const isMono = cs => cs.fontFamily.includes('Consolas') || cs.fontFamily.includes('monospace');

  function runsOf(el) {
    const items = [];
    (function collect(node) {
      if (node.nodeType === 3) {
        let t = node.textContent.replace(/[ \t\n\r]+/g, ' ');
        if (!t) return;
        const cs = getComputedStyle(node.parentElement);
        if (cs.textTransform === 'uppercase') t = t.toUpperCase();
        items.push({text: t, bold: parseInt(cs.fontWeight) >= 600, italic: cs.fontStyle === 'italic',
                    underline: cs.textDecorationLine.includes('underline'),
                    color: color(cs.color), mono: isMono(cs), size: px(cs.fontSize),
                    spacing: px(cs.letterSpacing)});
      } else if (node.nodeType === 1) {
        if (node.tagName === 'BR') items.push({br: true});
        else node.childNodes.forEach(collect);
      }
    })(el);
    // HTML sažima razmake: ukloni razmake na počecima i krajevima linija
    const lines = [[]];
    items.forEach(it => it.br ? lines.push([]) : lines[lines.length - 1].push(it));
    lines.forEach(line => {
      if (line.length) {
        line[0].text = line[0].text.replace(/^ +/, '');
        line[line.length - 1].text = line[line.length - 1].text.replace(/ +$/, '');
      }
    });
    return lines.map(line => line.filter(r => r.text.length));
  }

  // naslovi: prelomi redova tačno kao u browseru (inače PowerPoint ponekad stane u red manje)
  function headingLines(el, lines) {
    const out = [];
    let lastTop = null, line = [];
    // vrh svake reči u browseru; novi red počinje kad vrh skoči naniže
    const walker = document.createTreeWalker(el, NodeFilter.SHOW_TEXT);
    const words = [];
    let n;
    while ((n = walker.nextNode())) {
      const re = /\S+/g; let m;
      while ((m = re.exec(n.textContent))) {
        const r = document.createRange(); r.setStart(n, m.index); r.setEnd(n, m.index + m[0].length);
        const rect = r.getClientRects()[0];
        if (rect) words.push(rect.top);
      }
    }
    let wi = 0;
    lines.forEach((ln, li) => {
      if (li > 0) { out.push(line); line = []; lastTop = null; }
      ln.forEach(run => {
        let cur = '';
        (run.text.match(/\s*\S+\s*/g) || [run.text]).forEach(tok => {
          const top = words[wi++];
          if (lastTop !== null && top !== undefined && top > lastTop + 4) {
            if (cur) line.push({...run, text: cur.replace(/ +$/, '')});
            out.push(line); line = []; cur = tok.replace(/^ +/, '');
          } else cur += tok;
          if (top !== undefined) lastTop = top;
        });
        if (cur) line.push({...run, text: cur});
      });
    });
    out.push(line);
    return out.filter(l => l.length);
  }

  function block(el, cs) {
    const lh = cs.lineHeight === 'normal' ? px(cs.fontSize) * 1.2 : px(cs.lineHeight);
    return {align: cs.textAlign, lineHeight: lh, size: px(cs.fontSize),
            wrap: cs.whiteSpace !== 'nowrap', color: color(cs.color)};
  }

  const slides = [];
  document.querySelectorAll('section').forEach(sec => {
    const S = sec.getBoundingClientRect();
    const ops = [];
    const box = el => {
      const r = el.getBoundingClientRect();
      return {x: r.left - S.left, y: r.top - S.top, w: r.width, h: r.height};
    };

    function paint(el, cs, b) {
      const bg = color(cs.backgroundColor);
      const sides = ['Top', 'Right', 'Bottom', 'Left'].map(s => ({
        w: cs['border' + s + 'Style'] === 'none' ? 0 : px(cs['border' + s + 'Width']),
        style: cs['border' + s + 'Style'], color: color(cs['border' + s + 'Color'])}));
      const radii = ['TopLeft', 'TopRight', 'BottomRight', 'BottomLeft'].map(c => px(cs['border' + c + 'Radius']));
      const uniform = sides.every(s => s.w > 0 && s.w === sides[0].w && s.style === sides[0].style && s.color === sides[0].color);
      const radius = radii.every(r => r === radii[0]) ? radii[0] : 0;
      if (bg || uniform) {
        ops.push({t: 'rect', ...b, fill: bg, radius,
                  line: uniform ? {w: sides[0].w, color: sides[0].color, dash: sides[0].style === 'dashed'} : null});
      }
      if (!uniform) {
        const [T, R, B, L] = sides;
        if (T.w > 0 && T.color) ops.push({t: 'rect', x: b.x, y: b.y, w: b.w, h: T.w, fill: T.color, radius: 0});
        if (B.w > 0 && B.color) ops.push({t: 'rect', x: b.x, y: b.y + b.h - B.w, w: b.w, h: B.w, fill: B.color, radius: 0});
        if (L.w > 0 && L.color) ops.push({t: 'rect', x: b.x, y: b.y, w: L.w, h: b.h, fill: L.color, radius: 0});
        if (R.w > 0 && R.color) ops.push({t: 'rect', x: b.x + b.w - R.w, y: b.y, w: R.w, h: b.h, fill: R.color, radius: 0});
      }
    }

    // tekst ne sme da izađe iz najbližeg obojenog ili uokvirenog okvira
    function limits(el) {
      for (let a = el.parentElement; a && a !== sec; a = a.parentElement) {
        const cs = getComputedStyle(a);
        const painted = color(cs.backgroundColor) || px(cs.borderLeftWidth) > 0 || px(cs.borderRightWidth) > 0;
        if (!painted) continue;
        const r = a.getBoundingClientRect();
        return {limL: r.left - S.left + px(cs.borderLeftWidth) + px(cs.paddingLeft) * 0.4,
                limR: r.right - S.left - px(cs.borderRightWidth) - px(cs.paddingRight) * 0.4};
      }
      return {limL: 64, limR: 1920 - 64};
    }

    function content(cs, b) {
      const l = px(cs.paddingLeft) + px(cs.borderLeftWidth), r = px(cs.paddingRight) + px(cs.borderRightWidth);
      const t = px(cs.paddingTop) + px(cs.borderTopWidth), bo = px(cs.paddingBottom) + px(cs.borderBottomWidth);
      return {x: b.x + l, y: b.y + t, w: b.w - l - r, h: b.h - t - bo};
    }

    (function walk(el) {
      for (const ch of el.children) {
        if (ch.tagName === 'ASIDE') continue;
        const cs = getComputedStyle(ch);
        if (cs.display === 'none') continue;
        const b = box(ch);
        if (ch.tagName === 'X-ICON') { ops.push({t: 'icon', ...b, id: ch.getAttribute('data-icon-id')}); continue; }
        if (ch.tagName === 'X-SHAPE') {
          ops.push({t: 'shape', kind: ch.getAttribute('kind'), ...b, fill: color(cs.backgroundColor)});
          continue;
        }
        paint(ch, cs, b);
        if (TEXT.has(ch.tagName)) {
          let lines = runsOf(ch);
          if (/^H[123]$/.test(ch.tagName)) lines = headingLines(ch, lines);
          ops.push({t: 'text', ...content(cs, b), ...block(ch, cs), ...limits(ch), paras: [lines]});
          continue;
        }
        if (ch.tagName === 'UL' || ch.tagName === 'OL') {
          const items = [...ch.children].filter(li => li.tagName === 'LI').map(li => runsOf(li));
          ops.push({t: 'list', ordered: ch.tagName === 'OL', ...b, indent: px(cs.paddingLeft),
                    ...block(ch, cs), ...limits(ch), items});
          continue;
        }
        walk(ch);
      }
    })(sec);

    const aside = sec.querySelector(':scope > aside');
    slides.push({id: sec.id, background: color(getComputedStyle(sec).backgroundColor) || 'FFFFFF',
                 notes: aside ? aside.textContent.replace(/\s+/g, ' ').trim() : '', ops});
  });
  return slides;
}
"""


def load_deck():
    deck = json.loads((IZVOR / "deck.json").read_text(encoding="utf-8"))
    order = deck["order"]
    starts = [order.index(deck["sections"][f"p{i}"]["start"]) for i in range(1, 10)]
    bounds = [0] + starts[1:] + [len(order)]  # lekcija 1 uključuje i naslovne slajdove
    return [order[bounds[i]:bounds[i + 1]] for i in range(9)]


def page_html(slide_ids, with_web_fonts):
    sections = "\n".join(
        (IZVOR / "slides" / f"{sid}.html").read_text(encoding="utf-8") for sid in slide_ids)
    head = FONTS_LINK if with_web_fonts else ""
    return (f'<!doctype html><html lang="sr"><head><meta charset="utf-8">{head}'
            f"<style>{BASE_CSS}</style></head><body>{sections}</body></html>")


def emu(v):
    return Emu(int(round(v * EMU_PER_PX)))


def set_run(run, item):
    run.text = item["text"]
    font = run.font
    font.name = PPTX_MONO if item["mono"] else PPTX_SANS
    font.size = Pt(item["size"] * PT_PER_PX)
    font.bold = item["bold"]
    font.italic = item["italic"]
    if item["underline"]:
        font.underline = True
    if item["color"]:
        font.color.rgb = RGBColor.from_string(item["color"])
    if item["spacing"]:
        run._r.get_or_add_rPr().set("spc", str(int(round(item["spacing"] * PT_PER_PX * 100))))


def fill_paragraph(par, lines, op):
    par.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
                     "justify": PP_ALIGN.JUSTIFY}.get(op["align"], PP_ALIGN.LEFT)
    par.line_spacing = Pt(op["lineHeight"] * PT_PER_PX)
    par.space_before = Pt(0)
    par.space_after = Pt(0)
    for i, line in enumerate(lines):
        if i:
            par.add_line_break()
        for item in line:
            set_run(par.add_run(), item)
    # veličina fonta i za prazne linije i znak kraja pasusa
    end = par._p.get_or_add_endParaRPr()
    end.set("sz", str(int(round(op["size"] * PT_PER_PX * 100))))


def text_frame(slide, x, y, w, h, wrap):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(max(w, 1)), emu(max(h, 1)))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def widen(op):
    """Malo prostora viška, da PowerPoint ne prelomi red ranije nego browser.

    PowerPoint crta Arial i Consolas do ~1,5% šire od Chromiuma, pa je 2% dovoljno;
    okvir teksta ipak ne sme da izađe iz kartice ili obojenog polja u kome stoji.
    """
    x0, w0 = op["x"], op["w"]
    extra = max(4.0, w0 * 0.02)
    if op["align"] == "center":
        x, w = x0 - extra / 2, w0 + extra
    elif op["align"] == "right":
        x, w = x0 - extra, w0 + extra
    else:
        x, w = x0, w0 + extra
    left = max(min(op["limL"], x0), 0.0)
    right = min(max(op["limR"], x0 + w0), 1920.0)
    x = max(x, left)
    w = min(x + w, right) - x
    return x, w


def expected_lines_name(op):
    """Ime okvira beleži broj redova u browseru i razmak redova (za proveru u PowerPoint-u)."""
    lines = max(1, round(op["h"] / op["lineHeight"])) if op["lineHeight"] else 1
    return f"Tekst {lines}x{op['lineHeight'] * PT_PER_PX:.2f}"


def add_rect(slide, op):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if op["radius"] > 0 else MSO_SHAPE.RECTANGLE
    x, y, w, h = op["x"], op["y"], op["w"], op["h"]
    line = op.get("line")
    if line:  # PowerPoint crta liniju po sredini ivice, browser unutar okvira
        x, y, w, h = x + line["w"] / 2, y + line["w"] / 2, w - line["w"], h - line["w"]
    shp = slide.shapes.add_shape(kind, emu(x), emu(y), emu(max(w, 1)), emu(max(h, 1)))
    shp.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, op["radius"] / max(1.0, min(w, h)))
    if op["fill"]:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(op["fill"])
    else:
        shp.fill.background()
    if line and line["color"]:
        shp.line.color.rgb = RGBColor.from_string(line["color"])
        shp.line.width = Pt(line["w"] * PT_PER_PX)
        if line["dash"]:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shp.line.fill.background()
    return shp


def add_arrow(slide, op):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(op["x"]), emu(op["y"]), emu(op["w"]), emu(op["h"]))
    shp.shadow.inherit = False
    shp.adjustments[0] = 0.4  # debljina tela strelice, kao u browseru
    shp.adjustments[1] = 0.4 * op["w"] / max(1.0, op["h"])  # vrh zauzima poslednjih 40%
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(op["fill"] or "0F7C6E")
    shp.line.fill.background()


def add_bullets(par, op, ordered):
    pPr = par._p.get_or_add_pPr()
    indent = op["indent"]
    pPr.set("marL", str(int(round(indent * EMU_PER_PX))))
    pPr.set("indent", str(-int(round(indent * (1.0 if ordered else 0.75) * EMU_PER_PX))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    if op["color"]:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", op["color"])
    if ordered:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "+mj-lt")
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def build_pptx(slides, icon_png, out_path, title):
    prs = Presentation()
    prs.slide_width = Emu(1920 * EMU_PER_PX)
    prs.slide_height = Emu(1080 * EMU_PER_PX)
    blank = prs.slide_layouts[6]
    for sd in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(sd["background"])
        for op in sd["ops"]:
            t = op["t"]
            if t == "rect":
                add_rect(slide, op)
            elif t == "shape":
                if op["kind"] == "arrow-right":
                    add_arrow(slide, op)
            elif t == "icon":
                png = icon_png.get((sd["id"], op["id"]))
                if png:
                    slide.shapes.add_picture(png, emu(op["x"]), emu(op["y"]), emu(op["w"]), emu(op["h"]))
            elif t == "text":
                if not any(op["paras"][0]):
                    continue
                x, w = widen(op)
                box, tf = text_frame(slide, x, op["y"], w, op["h"], op["wrap"])
                fill_paragraph(tf.paragraphs[0], op["paras"][0], op)
                box.name = expected_lines_name(op)
            elif t == "list":
                x, w = widen(op)
                box, tf = text_frame(slide, x, op["y"], w, op["h"], op["wrap"])
                box.name = expected_lines_name(op)
                for i, item in enumerate(op["items"]):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    fill_paragraph(par, item, op)
                    add_bullets(par, op, op["ordered"])
        if sd["notes"]:
            slide.notes_slide.notes_text_frame.text = sd["notes"]
    props = prs.core_properties
    props.title = title
    props.author = AUTHOR
    props.subject = "Prema knjizi Martina Jonesa: Python for Biologists"
    prs.save(out_path)


def main(selected):
    sys.stdout.reconfigure(encoding="utf-8")
    lessons = load_deck()
    LEKCIJE.mkdir(exist_ok=True)
    report = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        for num, (slide_ids, (stem, name)) in enumerate(zip(lessons, LESSONS), start=1):
            if selected and num not in selected:
                continue
            title = f"{COURSE} – Lekcija {num}: {name}"

            # PDF, originalni fontovi
            page = browser.new_page(viewport={"width": 1920, "height": 1080})
            page.set_content(page_html(slide_ids, True), wait_until="networkidle")
            page.evaluate(ICON_SCRIPT, ICONS)
            page.evaluate("document.fonts.ready.then(() => true)")
            for item in page.evaluate(OVERFLOW_SCRIPT):
                report.append((num, item))
            page.pdf(path=str(LEKCIJE / f"{stem}.pdf"), width="1920px", height="1080px",
                     print_background=True, prefer_css_page_size=True)
            page.close()

            # PPTX, raspored meren sa Arial/Consolas
            page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=3)
            page.set_content(page_html(slide_ids, False), wait_until="load")
            page.evaluate(ICON_SCRIPT, ICONS)
            page.evaluate(MEASURE_FONTS_SCRIPT, [PPTX_SANS, PPTX_MONO])
            page.evaluate("document.fonts.ready.then(() => true)")
            slides = page.evaluate(EXTRACT_SCRIPT)
            icon_png = {}
            for sd in slides:
                for op in sd["ops"]:
                    if op["t"] == "icon":
                        loc = page.locator(f'section#{sd["id"]} [data-icon-id="{op["id"]}"]')
                        path = LEKCIJE.parent / "alati" / f".ikona-{sd['id']}-{op['id']}.png"
                        loc.screenshot(path=str(path), omit_background=True)
                        icon_png[(sd["id"], op["id"])] = str(path)
            page.close()
            build_pptx(slides, icon_png, LEKCIJE / f"{stem}.pptx", title)
            for p in icon_png.values():
                Path(p).unlink(missing_ok=True)
            print(f"Lekcija {num}: {len(slide_ids)} slajdova -> {stem}.pdf, {stem}.pptx")
        browser.close()
    if report:
        print("\nMoguća prekoračenja rasporeda:")
        for num, item in report:
            print(f"  L{num} {item['slide']} <{item['tag']}> {', '.join(item['issues'])}: {item['text']}")


if __name__ == "__main__":
    main({int(a) for a in sys.argv[1:]})
